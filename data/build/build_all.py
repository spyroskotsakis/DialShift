#!/usr/bin/env python3
"""Build the full radio catalog: canonical CSVs, the app catalog JSON and the multi-tab XLSX.

Usage:
    .venv/bin/python build/build_all.py [--refresh]

    --refresh   re-download raw sources (Wikipedia + radio-browser) for all countries

Outputs:
    canonical/<country>-stations.csv   canonical/collection-<code>.csv
    output/app-catalog.json            (the Add-station picker's catalog; validated, see app_catalog.py)
    output/dialshift-radio-catalog.xlsx

A bad languages.yaml (the app catalog's language table) exits non-zero before anything is written; a
failed app-catalog validation exits non-zero before the JSON and the XLSX are written.

The XLSX is deliberately plain (openpyxl basics only: freeze panes, auto-filter,
column widths) so it opens cleanly in macOS Numbers.
"""
import csv
import hashlib
import ssl
import sys
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app_catalog import build_app_catalog, unknown_language_name, validate_app_catalog, write_app_catalog
from build_stations import all_collections, all_countries, build_collection, build_country
from common import app_tag, language_table

DATA_DIR = Path(__file__).resolve().parent.parent
CANONICAL = DATA_DIR / 'canonical'
OUTPUT = DATA_DIR / 'output'
LANGUAGES = DATA_DIR / 'languages.yaml'         # the app catalog's language table (D84)
FAVICON_DIR = DATA_DIR / 'raw' / 'favicons'     # gitignored cache

CANON_COLS = ['country', 'name', 'name_local', 'city', 'region', 'frequency_fm', 'type', 'genre',
              'language', 'political_leaning', 'internet_only', 'stream_url', 'codec', 'bitrate',
              'stream_status', 'votes', 'logo', 'notes', 'source']

IMPORT_COLS = ['Station Name', 'Description / Genre', 'Stream URL', 'Country', 'City', 'Type',
               'Language', 'Political Leaning', 'Internet Only', 'Frequency FM', 'Votes', 'Notes',
               'Logo URL']

HEADER_FILL = PatternFill('solid', fgColor='2F5B8C')
HEADER_FONT = Font(bold=True, color='FFFFFF')


def write_sheet(wb, title, header, rows, widths=None, tab_color=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical='center')
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    return ws


_SSL_CTX = ssl._create_unverified_context()
_UA = {'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36'}


def _dl(url, timeout=6):
    req = urllib.request.Request(url, headers={**_UA, 'Accept': 'image/avif,image/webp,image/*,*/*;q=0.8',
                                               'Referer': 'https://www.google.com/'})
    with urllib.request.urlopen(req, timeout=timeout, context=_SSL_CTX) as r:
        return r.read(128 * 1024)


def fetch_logo(url, size=36):
    """Download a station logo, normalized to a small PNG (cached under raw/favicons/).
    Returns a local PNG path or None."""
    if not url:
        return None
    try:
        h = hashlib.md5(url.encode()).hexdigest()
        png = FAVICON_DIR / f'{h}.png'
        if png.exists():
            return png
        candidates = [url]
        # Wikimedia serves SVG thumbnails; the same path with '.png' appended is
        # the rasterized render (Pillow cannot decode SVG).
        if 'wikimedia' in url.lower() and url.lower().rstrip().endswith('.svg'):
            candidates.append(url + '.png')
        if url.startswith('http://'):
            candidates.append('https://' + url[7:])
        elif url.startswith('https://'):
            candidates.append('http://' + url[8:])
        data = None
        for c in candidates:
            try:
                data = _dl(c)
                break
            except Exception:
                continue
        if not data:
            return None
        from io import BytesIO
        from PIL import Image
        img = Image.open(BytesIO(data))
        img = img.convert('RGBA')
        img.thumbnail((size, size), Image.Resampling.LANCZOS)
        FAVICON_DIR.mkdir(parents=True, exist_ok=True)
        img.save(png, 'PNG')
        return png
    except Exception:
        return None


def embed_logos(ws, logo_col, logo_paths):
    """Embed 18px logo images into the given column, one per data row.
    logo_paths: list of (row_number, png_path). Rows without a logo stay text-only."""
    from openpyxl.drawing.image import Image as _Img
    for r, path in logo_paths:
        try:
            img = _Img(str(path))
            img.width = 18
            img.height = 18
            ws.add_image(img, f'{logo_col}{r}')
            ws.row_dimensions[r].height = 19
        except Exception:
            continue


def load_languages():
    """data/languages.yaml checked by common.language_table; a bad table stops the run before anything is
    built or written."""
    try:
        with open(LANGUAGES, encoding='utf-8') as f:
            return language_table(yaml.safe_load(f), LANGUAGES.name)
    except ValueError as e:                        # yaml.YAMLError is not a ValueError: it keeps its traceback
        print(f'languages: {e}', file=sys.stderr)
        sys.exit(1)


def main():
    refresh = '--refresh' in sys.argv
    languages = load_languages()
    countries = all_countries()
    per_country = {}
    for cfg in countries:
        rows, stats = build_country(cfg, force_refresh=refresh)
        per_country[cfg['code']] = rows
        CANONICAL.mkdir(parents=True, exist_ok=True)
        csv_path = CANONICAL / f"{cfg['name'].lower().replace(' ', '-')}-stations.csv"
        with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            wr = csv.DictWriter(f, fieldnames=CANON_COLS, extrasaction='ignore')
            wr.writeheader()
            for r in rows:
                wr.writerow({k: r.get(k, '') for k in CANON_COLS})
        print(f"{cfg['code']} {cfg['name']}: {len(rows)} rows -> {csv_path.name}")

    collections = all_collections()
    per_collection = {}
    for cfg in collections:
        rows = build_collection(cfg)
        per_collection[cfg['name']] = rows
        csv_path = CANONICAL / f"collection-{cfg['code'].lower()}.csv"
        with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            wr = csv.DictWriter(f, fieldnames=CANON_COLS, extrasaction='ignore')
            wr.writeheader()
            for r in rows:
                wr.writerow({k: r.get(k, '') for k in CANON_COLS})
        print(f"[collection] {cfg['name']}: {len(rows)} rows -> {csv_path.name}")

    # ---------------------------------------------------------------- app catalog (JSON)
    sources = [(cfg, per_country[cfg['code']]) for cfg in countries] + \
              [(cfg, per_collection[cfg['name']]) for cfg in collections]
    app_doc, app_stats = build_app_catalog(
        sources, datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'), languages)
    app_json = OUTPUT / 'app-catalog.json'
    app_json_rel = app_json.relative_to(DATA_DIR.parent)
    print(f"app-catalog: working={app_stats['working']} url_excluded={app_stats['url_excluded']} "
          f"duplicates_removed={app_stats['duplicates_removed']} exported={app_stats['exported']} "
          f"-> {app_json_rel}")
    unknown = app_stats['unknown_languages']
    print(f"app-catalog: languages={app_stats['languages']} unknown={len(unknown)}")
    for _, (first, count) in sorted(unknown.items(), key=lambda kv: (-kv[1][1], kv[0])):
        print(f'app-catalog: unknown language "{first}" in {count} entries, exported as '
              f'"{unknown_language_name(first)}"; add it to data/{LANGUAGES.name}')
    problems = validate_app_catalog(app_doc, [r for _, rows in sources for r in rows], languages)
    if problems:
        print(f'app-catalog: validation failed with {len(problems)} problem(s); {app_json_rel} is unchanged '
              f'(fix the YAML, not the script):', file=sys.stderr)
        for problem in problems:
            print(f'  - {problem}', file=sys.stderr)
        sys.exit(1)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    write_app_catalog(app_doc, app_json)

    all_rows = [r for rows in per_country.values() for r in rows]
    working = [r for r in all_rows if r['stream_status'] == 'Working' and r['stream_url']]

    # ---------------------------------------------------------------- workbook
    wb = Workbook()
    wb.remove(wb.active)

    # ---- README tab ----
    ws = wb.create_sheet('README')
    ws.sheet_properties.tabColor = 'C0C0C0'
    ws.column_dimensions['A'].width = 118
    ws.column_dimensions['B'].width = 46
    ws.column_dimensions['C'].width = 46
    bold, italic = Font(bold=True, size=14), Font(italic=True, color='808080')
    section = Font(bold=True, size=12, color='FFFFFF')
    section_fill = PatternFill('solid', fgColor='2F5B8C')

    def line(text, font=None, fill=None):
        ws.append([text, '', ''])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if font: c.font = font
            if fill: c.fill = fill

    def section_row(text):
        line(text, font=section, fill=section_fill)

    line('DialShift Radio Catalog', font=bold)
    line('Generated by data/build/build_all.py — never edit this workbook by hand. '
         'Edit the country files in data/countries/*.yaml and rebuild.', font=italic)
    line('')

    section_row('How to add a station to the DialShift app')
    line('The app has this catalog built in: every station with a working stream, all countries and '
         f'collections ({app_stats["exported"]} stations in this build).')
    line('1.  Open the app → Stations → Add a frequency.')
    line('2.  Type in the search box: a station name, a city or an FM frequency (1015 or 101.5). '
         'Narrow the list with the Country, City, Type, Genre and Language filters if you like.')
    line('3.  Pick a result (click it, or Down and Enter). It fills Station name, Description / genre '
         'and Stream URL, and shows the station notes. Change any field if you want, then Save.')
    line('Fallback — manual entry, for a station that is not in the catalog or when the app shows '
         '"Catalog unavailable": open the "Import Ready" tab and copy the three fields into the form under '
         '"Or enter stream details manually":')
    line('       Station name        ←  "Station Name" column')
    line('       Description / genre ←  "Description / Genre" column')
    line('       Stream URL          ←  "Stream URL" column')
    line('The app plays MP3, AAC and HLS streams. Webpage URLs do not play.')
    line('')

    section_row('Tabs in this workbook')
    tab_descs = [
        ('Import Ready', 'every station with a working stream, all countries — copy from here when '
                         'you enter a station by hand'),
        (' / '.join(cfg['name'] for cfg in countries),
         'full per-country lists, including stations without a working stream'),
    ]
    for cfg in countries:
        for farea in cfg.get('focus_areas', []):
            place = farea.get('city') or farea.get('region')
            tab_descs.append((farea['label'],
                              f"your local shortlist for {cfg['name']}: {place} stations"))
    for cfg in collections:
        tab_descs.append((cfg['name'], cfg.get('description', '').replace('\n', ' ')))
    tab_descs.append(('Summary', 'counts per country and type, top cities'))
    for title, desc in tab_descs:
        line(f'  •  {title}:  {desc}')
    line('')

    section_row('Stream status & how to refresh the data')
    line('  stream_status comes from radio-browser.info checks (Working / Down / No stream found).')
    line('  To refresh everything:  data/.venv/bin/python data/build/build_all.py --refresh')
    line('  The same run regenerates the app\'s catalog (data/output/app-catalog.json) and this workbook.')
    line('')

    # glossary table
    ws.append(['Column', 'Meaning', ''])
    for c in ws[ws.max_row]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    glossary = [
        ('type', 'Music | News & Talk | Political | Sports | Religious | Municipal | Public | Other'),
        ('genre', 'music style or programme focus (laïkó, rock, classical, news, sports…)'),
        ('language', 'main broadcast language (Greek / French / German / English / Multilingual)'),
        ('political_leaning', 'Left / Center / Right / State / Municipal — only where well documented; None = not applicable'),
        ('internet_only', 'Yes = web-only · No = terrestrial (FM/AM/DAB) · Unknown = not listed in an official FM directory'),
        ('frequency_fm', 'FM frequency where known (blank for internet-only)'),
        ('votes', 'popularity signal from radio-browser.info — higher means more listeners there'),
        ('logo / Logo URL', 'station logo: focus tabs and collections show the image; Logo URL is the source link'),
        ('notes', 'description from Wikipedia or curated notes'),
        ('source', 'where the row came from: curated facts, Wikipedia, radio-browser'),
    ]
    for name, meaning in glossary:
        ws.append([name, meaning, ''])
        ws[f'A{ws.max_row}'].font = Font(bold=True)
        ws[f'B{ws.max_row}'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.append(['', '', ''])
    line('Station names and streams belong to their respective owners. '
         'DialShift connects directly to each provider; this catalog is unaffiliated.', font=italic)

    # ---- Summary tab ----
    summary = [['Country', 'Stations', 'Working streams', 'Internet-only', 'Focus tabs']]
    for cfg in countries:
        rows = per_country[cfg['code']]
        labels = ', '.join(f['label'] for f in cfg.get('focus_areas', []))
        summary.append([cfg['name'], len(rows), sum(1 for r in rows if r['stream_status'] == 'Working'),
                        sum(1 for r in rows if r['internet_only'] == 'Yes'), labels])
    summary.append([])
    summary.append(['Focus tabs detail'])
    for cfg in countries:
        for farea in cfg.get('focus_areas', []):
            n = sum(1 for r in per_country[cfg['code']] if r.get('focus_area') == farea['label'])
            summary.append([f"{cfg['name']} · {farea['label']}", n])
    summary.append([])
    summary.append(['Collections (internet radio)'])
    for cfg in collections:
        rows = per_collection[cfg['name']]
        summary.append([cfg['name'], len(rows), sum(1 for r in rows if r['stream_status'] == 'Working'),
                        'Yes (all)', ''])
    summary.append([])
    summary.append(['By type (all countries)'])
    for t, n in Counter(r['type'] for r in all_rows).most_common():
        summary.append([t, n])
    summary.append([])
    summary.append(['Top cities (all countries)'])
    for c, n in Counter(r['city'] for r in all_rows if r['city'] not in ('', '—')).most_common(15):
        summary.append([c, n])
    write_sheet(wb, 'Summary', ['Item', 'Count'], summary, widths=[40, 18], tab_color='FFE699')

    # ---- Import Ready tab ----
    country_names = {cfg['code']: cfg['name'] for cfg in countries}
    imp = []
    for r in sorted(working, key=lambda x: (x['country'], -(x['votes'] or 0))):
        imp.append([r['name'], app_tag(r), r['stream_url'],
                    country_names.get(r['country'], r['country']),
                    r['city'], r['type'], r['language'], r['political_leaning'],
                    r['internet_only'], r['frequency_fm'], r['votes'] or '', r['notes'][:180],
                    r.get('logo', '')])
    write_sheet(wb, 'Import Ready', IMPORT_COLS, imp,
                widths=[34, 30, 58, 10, 20, 14, 12, 16, 12, 10, 9, 40, 40], tab_color='90EE90')

    # ---- per-country tabs ----
    XLS_COLS = ['Name', 'Name (local)', 'City', 'Region', 'Frequency FM', 'Type', 'Genre', 'Language',
                'Political Leaning', 'Internet Only', 'Stream URL', 'Codec', 'Bitrate', 'Stream Status',
                'Votes', 'Notes', 'Source', 'Logo URL']
    for cfg in countries:
        rows = per_country[cfg['code']]
        data = [[r['name'], r['name_local'], r['city'], r['region'], r['frequency_fm'], r['type'],
                 r['genre'], r['language'], r['political_leaning'], r['internet_only'], r['stream_url'],
                 r['codec'], r['bitrate'], r['stream_status'], r['votes'] or '', r['notes'][:200],
                 r['source'], r.get('logo', '')]
                for r in rows]
        write_sheet(wb, cfg['name'], XLS_COLS, data,
                    widths=[32, 24, 20, 16, 10, 13, 20, 11, 16, 11, 55, 7, 8, 13, 8, 40, 16, 40],
                    tab_color='9DC3E6')

    # ---- focus tabs (Munich, Paris, Toulouse, Aude, …) ----
    for cfg in countries:
        for farea in cfg.get('focus_areas', []):
            label = farea['label']
            # curated focus stations + terrestrial stations in the focus city/region
            focus_rows = [r for r in per_country[cfg['code']] if r.get('focus_area') == label]
            # dedupe by (name, url)
            seen = set()
            unique = []
            for r in focus_rows:
                k = (r['name'].lower(), r['stream_url'])
                if k not in seen:
                    seen.add(k)
                    unique.append(r)
            unique = sorted(unique, key=lambda x: (x['stream_url'] == '', -x['votes'], x['name'].lower()))
            data = [[r['name'], app_tag(r), r['stream_url'], r['city'], r['frequency_fm'], r['type'],
                     r['language'], r['stream_status'], r['votes'] or '', r['notes'][:160]]
                    for r in unique]
            ws = write_sheet(wb, label,
                             ['Station Name', 'Logo', 'Description / Genre', 'Stream URL', 'City',
                              'Frequency FM', 'Type', 'Language', 'Stream Status', 'Votes', 'Notes',
                              'Logo URL'],
                             [[row[0], '', row[1], row[2], row[3], row[4], row[5], row[6], row[7],
                               row[8], row[9], ''] for row in data],
                             widths=[34, 5, 30, 58, 20, 10, 14, 11, 13, 9, 40, 40], tab_color='F4B183')
            # download + embed logos for this tab (threaded)
            logo_urls = [r.get('logo', '') for r in unique]
            with ThreadPoolExecutor(max_workers=8) as ex:
                paths = list(ex.map(fetch_logo, logo_urls))
            embed_logos(ws, 'B', [(i + 2, p) for i, p in enumerate(paths) if p])
            # Logo URL column
            for i, r in enumerate(unique):
                ws.cell(row=i + 2, column=12).value = r.get('logo', '')
            print(f"{label} tab: {len(data)} stations, "
                  f"{sum(1 for p in paths if p)} logos embedded")

    # ---- collection tabs (genre folders) ----
    for cfg in collections:
        rows = per_collection[cfg['name']]
        data = [[r['name'], app_tag(r), r['stream_url'], r['language'],
                 r['stream_status'], r['votes'] or '', r['notes'][:160]]
                for r in rows]
        ws = write_sheet(wb, cfg['name'],
                         ['Station Name', 'Logo', 'Description / Genre', 'Stream URL', 'Language',
                          'Stream Status', 'Votes', 'Notes', 'Logo URL'],
                         [[row[0], '', row[1], row[2], row[3], row[4], row[5], row[6], '']
                          for row in data],
                         widths=[30, 5, 34, 58, 13, 13, 9, 40, 40], tab_color='C9AED6')
        logo_urls = [r.get('logo', '') for r in rows]
        with ThreadPoolExecutor(max_workers=8) as ex:
            paths = list(ex.map(fetch_logo, logo_urls))
        embed_logos(ws, 'B', [(i + 2, p) for i, p in enumerate(paths) if p])
        for i, r in enumerate(rows):
            ws.cell(row=i + 2, column=9).value = r.get('logo', '')
        print(f"{cfg['name']} tab: {len(data)} stations, "
              f"{sum(1 for p in paths if p)} logos embedded")

    OUTPUT.mkdir(parents=True, exist_ok=True)
    xlsx = OUTPUT / 'dialshift-radio-catalog.xlsx'
    wb.save(xlsx)
    print(f'\nwrote {xlsx}')
    print(f'  tabs: {wb.sheetnames}')
    print(f'  total stations: {len(all_rows)} | working streams: {len(working)}')


if __name__ == '__main__':
    main()
